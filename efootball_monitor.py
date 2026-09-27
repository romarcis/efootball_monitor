#!/usr/bin/env python3
"""
eFootball Wi-Fi Monitor
=======================

Monitora la connessione mentre giochi a eFootball su PC (Windows, Wi-Fi).

Ogni secondo misura:
  * ping / jitter / perdita verso il ROUTER (qualita' del Wi-Fi di casa)
  * ping / jitter / perdita verso INTERNET (default 1.1.1.1, qualita' della linea)
  * segnale Wi-Fi (%, RSSI, canale, banda, velocita' di collegamento)

Con l'opzione --sniff (richiede Npcap + scapy) osserva anche il traffico UDP
della partita, come fa eFootball_Network_Monitor_Tool:
  * individua il server / avversario della partita (IP e porta)
  * pacchetti al secondo ricevuti e inviati
  * "scatti": pause nei pacchetti in arrivo (> 150 ms) = lag lato server->te
  * ping verso il server della partita (se risponde all'ICMP)

Ogni intervallo (default 5 s) stampa una riga colorata.
Premi Ctrl+C per fermare: stampa il riepilogo e crea il file Excel della sessione,
piu' lo storico efootball_match_history.xlsx (con --sniff).

Servono: pip install openpyxl (Excel) e, per --sniff, Npcap + pip install scapy.
"""

import argparse
import csv
import ctypes
import datetime as dt
import ipaddress
import json
import os
import platform
import re
import socket
import statistics
import struct
import subprocess
import sys
import threading
import time
import urllib.request
from collections import defaultdict, deque

IS_WINDOWS = platform.system() == "Windows"
FROZEN = getattr(sys, "frozen", False)  # True quando gira come .exe (PyInstaller)

# --------------------------------------------------------------------------
# Lingua: italiano o inglese (--lang it|en, default: lingua di Windows)
# --------------------------------------------------------------------------
LANG = "it"


def T(it, en):
    """Testo nella lingua scelta."""
    return it if LANG == "it" else en


def detect_lang():
    try:
        if IS_WINDOWS:
            # LANG_ITALIAN = 0x10 nella parte bassa dell'identificativo di lingua
            if ctypes.windll.kernel32.GetUserDefaultUILanguage() & 0x3FF == 0x10:
                return "it"
            return "en"
        import locale
        return "it" if (locale.getlocale()[0] or "").lower().startswith("it") else "en"
    except Exception:
        return "en"


def set_lang(lang):
    """Imposta la lingua e tutte le etichette che finiscono nei file."""
    global LANG, OK, WARN, BAD, CONN_GOOD, CONN_SOSO, CONN_NA, CONN_POOR, YES, NO
    global KIND_SERVER, KIND_P2P, CSV_FIELDS, HISTORY_HEADER
    LANG = lang
    OK, WARN, BAD = "OK", T("ATTENZIONE", "WARNING"), T("PROBLEMA", "PROBLEM")
    CONN_GOOD, CONN_SOSO = T("BUONA", "GOOD"), T("COSI' COSI'", "SO-SO")
    CONN_NA, CONN_POOR = T("NON MISURABILE", "NOT MEASURABLE"), T("SCARSA", "POOR")
    YES, NO = T("SI", "YES"), "no"
    KIND_SERVER = T("server dedicato", "dedicated server")
    KIND_P2P = T("P2P (diretta con l'avversario)", "P2P (direct to opponent)")
    CSV_FIELDS = [T(it, en) for it, en in _CSV_FIELDS]
    HISTORY_HEADER = [T(it, en) for it, en in _HISTORY_HEADER]


# Etichette dei giudizi in entrambe le lingue -> livello (0 ok, 1 attenzione, 2 problema)
LEVEL_OF = {"OK": 0, "ATTENZIONE": 1, "WARNING": 1, "PROBLEMA": 2, "PROBLEM": 2}
CONN_LEVEL = {"BUONA": 0, "GOOD": 0, "COSI' COSI'": 1, "SO-SO": 1, "NON MISURABILE": 1,
              "NOT MEASURABLE": 1, "SCARSA": 2, "POOR": 2}

# --------------------------------------------------------------------------
# Colori terminale
# --------------------------------------------------------------------------
if IS_WINDOWS:
    os.system("")  # abilita le sequenze ANSI nel prompt di Windows 10/11

GREEN, YELLOW, RED, CYAN, DIM, RESET = (
    "\033[92m", "\033[93m", "\033[91m", "\033[96m", "\033[2m", "\033[0m")


def color(value, text, good, bad):
    """Verde se value <= good, rosso se >= bad, giallo in mezzo."""
    if value is None:
        return f"{DIM}{text}{RESET}"
    if value <= good:
        return f"{GREEN}{text}{RESET}"
    if value >= bad:
        return f"{RED}{text}{RESET}"
    return f"{YELLOW}{text}{RESET}"


# --------------------------------------------------------------------------
# Ping
# --------------------------------------------------------------------------
class _IcmpPinger:
    """Ping via API IcmpSendEcho di Windows: niente processi, niente
    problemi di lingua nell'output di ping.exe, precisione al millisecondo."""

    class _Reply(ctypes.Structure):
        _fields_ = [("Address", ctypes.c_ulong), ("Status", ctypes.c_ulong),
                    ("RoundTripTime", ctypes.c_ulong), ("DataSize", ctypes.c_ushort),
                    ("Reserved", ctypes.c_ushort), ("Data", ctypes.c_void_p),
                    ("Ttl", ctypes.c_ubyte), ("Tos", ctypes.c_ubyte),
                    ("Flags", ctypes.c_ubyte), ("OptionsSize", ctypes.c_ubyte),
                    ("OptionsData", ctypes.c_void_p)]

    def __init__(self):
        self.lib = ctypes.windll.iphlpapi
        self.lib.IcmpCreateFile.restype = ctypes.c_void_p
        self.lib.IcmpSendEcho.argtypes = [
            ctypes.c_void_p, ctypes.c_ulong, ctypes.c_void_p, ctypes.c_ushort,
            ctypes.c_void_p, ctypes.c_void_p, ctypes.c_ulong, ctypes.c_ulong]
        self.lib.IcmpCloseHandle.argtypes = [ctypes.c_void_p]
        self.handle = self.lib.IcmpCreateFile()

    def ping(self, ip, timeout_ms):
        addr = struct.unpack("<L", socket.inet_aton(ip))[0]
        data = b"efootball-monitor"
        buf = ctypes.create_string_buffer(ctypes.sizeof(self._Reply) + len(data) + 8)
        n = self.lib.IcmpSendEcho(self.handle, addr, data, len(data), None,
                                  buf, len(buf), timeout_ms)
        if n == 0:
            return None
        reply = self._Reply.from_buffer(buf)
        if reply.Status != 0:  # 0 = IP_SUCCESS
            return None
        return float(reply.RoundTripTime)


def _ping_subprocess(ip, timeout_ms):
    """Fallback per Linux/macOS (utile solo per provare lo script)."""
    try:
        out = subprocess.run(["ping", "-c", "1", "-W", str(max(1, timeout_ms // 1000)), ip],
                             capture_output=True, text=True, timeout=timeout_ms / 1000 + 2)
    except Exception:
        return None
    m = re.search(r"time[=<]\s*([\d.]+)\s*ms", out.stdout)
    return float(m.group(1)) if m else None


def make_ping_fn():
    if IS_WINDOWS:
        try:
            pinger = _IcmpPinger()
            return pinger.ping
        except Exception:
            pass
    return _ping_subprocess


class PingTarget:
    """Pinga un host a intervallo fisso in un thread separato e
    accumula i risultati per la finestra corrente."""

    def __init__(self, name, ip, interval, timeout_ms, ping_fn, may_block_icmp=False):
        self.name, self.ip = name, ip
        self.may_block_icmp = may_block_icmp  # i server di gioco spesso ignorano il ping
        self.interval, self.timeout_ms, self.ping_fn = interval, timeout_ms, ping_fn
        self.lock = threading.Lock()
        self.samples = []            # rtt (float) o None se perso
        self.ever_replied = False
        self.all_rtts, self.all_sent, self.all_lost = [], 0, 0
        self._stop = threading.Event()
        threading.Thread(target=self._run, daemon=True).start()

    def set_ip(self, ip):
        with self.lock:
            if ip != self.ip:
                self.ip, self.samples, self.ever_replied = ip, [], False

    def _run(self):
        while not self._stop.is_set():
            start = time.monotonic()
            ip = self.ip
            if ip:
                rtt = self.ping_fn(ip, self.timeout_ms)
                with self.lock:
                    if ip == self.ip:
                        self.samples.append(rtt)
                        if rtt is not None:
                            self.ever_replied = True
            self._stop.wait(max(0.0, self.interval - (time.monotonic() - start)))

    def take_window(self):
        """Restituisce le statistiche della finestra e la azzera."""
        with self.lock:
            samples, self.samples = self.samples, []
            ever = self.ever_replied
        sent = len(samples)
        rtts = [s for s in samples if s is not None]
        lost = sent - len(rtts)
        self.all_sent += sent
        self.all_lost += lost
        self.all_rtts += rtts
        res = {"sent": sent, "lost": lost, "avg": None, "max": None,
               "jitter": None, "loss": None,
               "no_icmp": self.may_block_icmp and sent > 0 and not ever}
        if sent:
            res["loss"] = 100.0 * lost / sent
        if rtts:
            res["avg"] = statistics.fmean(rtts)
            res["max"] = max(rtts)
            # jitter = media delle differenze tra ping consecutivi (come RFC 3550)
            if len(rtts) > 1:
                res["jitter"] = statistics.fmean(abs(a - b) for a, b in zip(rtts, rtts[1:]))
            else:
                res["jitter"] = 0.0
        return res

    def stop(self):
        self._stop.set()


# --------------------------------------------------------------------------
# Rete locale: gateway, IP locale, Wi-Fi
# --------------------------------------------------------------------------
def local_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("1.1.1.1", 80))  # nessun pacchetto inviato, serve solo a scegliere l'interfaccia
        return s.getsockname()[0]
    except OSError:
        return None
    finally:
        s.close()


def default_gateway():
    try:
        if IS_WINDOWS:
            out = subprocess.run(["route", "print", "-4", "0.0.0.0"], capture_output=True,
                                 text=True, errors="replace").stdout
            best = None
            for line in out.splitlines():
                p = line.split()
                if len(p) >= 5 and p[0] == "0.0.0.0" and p[1] == "0.0.0.0":
                    try:
                        ipaddress.IPv4Address(p[2])
                        metric = int(p[4])
                    except ValueError:
                        continue
                    if best is None or metric < best[1]:
                        best = (p[2], metric)
            return best[0] if best else None
        out = subprocess.run(["ip", "route", "show", "default"], capture_output=True,
                             text=True).stdout
        m = re.search(r"default via (\S+)", out)
        return m.group(1) if m else None
    except Exception:
        return None


# Etichette di "netsh wlan show interfaces" in inglese e italiano
# Etichette di "netsh wlan show interfaces" in inglese e italiano.
# Alcune versioni di Windows 11 stampano righe spezzate ("...8d:57Banda\n : 5 GHz"),
# quindi i campi si cercano sul testo con gli spazi e gli a capo compattati.
_WLAN_FIELDS = {
    "bssid": r"BSSID(?:\s+AP)?\s*:\s*([0-9a-f]{2}(?::[0-9a-f]{2}){5})",
    "radio": r"(?:Radio type|Tipo (?:frequenza )?radio)\s*:\s*(802\.11\w+)",
    "band": r"(?:Band|Banda)\s*:\s*([\d.,]+\s*GHz)",
    "channel": r"(?:Channel|Canale)\s*:\s*(\d+)",
    "rx_mbps": r"(?:Receive rate|Velocit\S*\s+(?:di\s+)?ricezione)[^:]*:\s*([\d.,]+)",
    "tx_mbps": r"(?:Transmit rate|Velocit\S*\s+(?:di\s+)?trasmissione)[^:]*:\s*([\d.,]+)",
    "signal": r"(?:Signal|Segnale)\s*:\s*(\d+)\s*%",
    "rssi": r"Rssi\s*:\s*(-?\d+)",
}


_LOCATION_WARNED = False


def wifi_info():
    if not IS_WINDOWS:
        return {}
    try:
        raw = subprocess.run(["netsh", "wlan", "show", "interfaces"],
                             capture_output=True).stdout
        text = raw.decode("cp850", errors="replace")
    except Exception:
        return {}
    info = {}
    if "SSID" not in text and re.search(r"posizione|location", text, re.IGNORECASE):
        # Windows 11 24H2: netsh wlan funziona solo con i servizi di posizione attivi
        global _LOCATION_WARNED
        if not _LOCATION_WARNED:
            _LOCATION_WARNED = True
            print(YELLOW + T("Windows blocca la lettura del segnale Wi-Fi perche' la posizione "
                             "e' disattivata. Riattivala in Impostazioni > Privacy e "
                             "sicurezza > Posizione per vedere segnale, banda e canale.",
                             "Windows blocks reading the Wi-Fi signal because Location is "
                             "turned off. Turn it on in Settings > Privacy & security > "
                             "Location to see signal, band and channel.") + RESET)
        return info
    m = re.search(r"^\s*SSID\s*:\s*(.+?)\s*$", text, re.MULTILINE)
    if m:
        info["ssid"] = m.group(1)
    flat = re.sub(r"\s+", " ", text)
    for key, pat in _WLAN_FIELDS.items():
        m = re.search(pat, flat, re.IGNORECASE)
        if m:
            info[key] = m.group(1).strip().replace(",", ".")
    for k in ("signal", "rssi", "channel"):
        if k in info:
            info[k] = int(info[k])
    for k in ("rx_mbps", "tx_mbps"):
        if k in info:
            try:
                info[k] = float(info[k])
            except ValueError:
                del info[k]
    return info


# --------------------------------------------------------------------------
# Sniffer UDP (opzionale, stesso principio di eFootball_Network_Monitor_Tool)
# --------------------------------------------------------------------------
IGNORED_PORTS = {53, 67, 68, 123, 137, 138, 443, 1900, 3478, 3702, 5353, 5355}
GAP_MS = 150  # pausa oltre la quale consideriamo un "buco" nei pacchetti in arrivo


class MatchSniffer:
    """Conta i pacchetti UDP per host remoto. L'host con piu' traffico
    costante (> min_pps) e' considerato il server/avversario della partita."""

    def __init__(self, my_ip, iface=None, min_pps=15, on_new_remote=None):
        from scapy.all import AsyncSniffer, IP, UDP  # noqa: F401  (import pigro)
        self._IP, self._UDP = IP, UDP
        self.my_ip, self.min_pps = my_ip, min_pps
        # avviso anticipato: host nuovo che manda >= 5 pacchetti in 1 s
        self.on_new_remote = on_new_remote
        self.burst = defaultdict(deque)
        self.last_seen = {}
        self.ttl = {}                       # ultimo TTL ricevuto da ogni host
        self.lock = threading.Lock()
        self.rx = defaultdict(int)
        self.tx = defaultdict(int)
        self.ports = {}
        self.last_rx = {}
        self.gaps = defaultdict(list)
        self.win_first = {}                 # primo pacchetto ricevuto nella finestra
        self.rates = defaultdict(lambda: deque(maxlen=120))  # pacchetti/s osservati
        self.current = None
        self.sniffer = AsyncSniffer(filter=f"udp and host {my_ip}", prn=self._on_pkt,
                                    store=False, iface=iface)
        self.sniffer.start()

    def _on_pkt(self, pkt):
        IP, UDP = self._IP, self._UDP
        if IP not in pkt or UDP not in pkt:
            return
        ip, udp = pkt[IP], pkt[UDP]
        now = time.monotonic()
        with self.lock:
            if ip.dst == self.my_ip:
                remote, rport = ip.src, udp.sport
                if rport in IGNORED_PORTS or udp.dport in IGNORED_PORTS:
                    return
                self.rx[remote] += 1
                self.ports[remote] = rport
                last = self.last_rx.get(remote)
                if last is not None:
                    self.gaps[remote].append((now - last) * 1000)
                self.last_rx[remote] = now
                self.win_first.setdefault(remote, now)
                self.ttl[remote] = ip.ttl
                if self.on_new_remote and not ipaddress.ip_address(remote).is_private:
                    self._check_new(remote, rport, now)
            elif ip.src == self.my_ip:
                remote = ip.dst
                if udp.dport in IGNORED_PORTS or udp.sport in IGNORED_PORTS:
                    return
                self.tx[remote] += 1

    def _check_new(self, remote, rport, now):
        prev = self.last_seen.get(remote)
        self.last_seen[remote] = now
        if prev is not None and now - prev < 60 and remote not in self.burst:
            return  # host gia' attivo (o gia' annunciato): non e' una partita nuova
        b = self.burst[remote]
        b.append(now)
        while b and now - b[0] > 1.0:
            b.popleft()
        if len(b) >= 5:
            del self.burst[remote]
            threading.Thread(target=self.on_new_remote,
                             args=(remote, rport, self.ttl.get(remote)), daemon=True).start()

    def take_window(self, seconds):
        with self.lock:
            rx, tx, gaps = self.rx, self.tx, self.gaps
            self.rx, self.tx, self.gaps = defaultdict(int), defaultdict(int), defaultdict(list)
            first, self.win_first = self.win_first, {}
            last = dict(self.last_rx)
            ports = dict(self.ports)
        candidates = [(n, r) for r, n in rx.items()
                      if not ipaddress.ip_address(r).is_private and n / seconds >= self.min_pps]
        if not candidates:
            self.current = None
            return None
        _, remote = max(candidates)
        new = remote != self.current
        self.current = remote
        g = gaps.get(remote, [])
        n = rx[remote]
        return {
            "ip": remote, "port": ports.get(remote), "new": new,
            "rx_pps": n / seconds, "tx_pps": tx.get(remote, 0) / seconds,
            "max_gap": max(g) if g else None,
            "gaps": sum(1 for x in g if x > GAP_MS),
            "loss": self._estimate_loss(remote, n, first.get(remote), last.get(remote)),
            "rx": n,
        }

    def _estimate_loss(self, remote, n, t_first, t_last):
        """Stima dei pacchetti persi server -> te. Il server manda a ritmo fisso
        (circa 55 al secondo): il ritmo "pieno" e' quello delle finestre migliori,
        e quello che manca rispetto al tempo in cui la partita era attiva e' perso.
        Contare solo tra primo e ultimo pacchetto evita falsi allarmi a inizio/fine partita."""
        if t_first is None or t_last is None or n < 10:
            return None
        span = t_last - t_first
        if span < 0.5:
            return None
        rates = self.rates[remote]
        rates.append((n - 1) / span)
        nominal = sorted(rates)[int(len(rates) * 0.9)] if len(rates) >= 3 else max(rates)
        expected = span * nominal + 1
        return max(0.0, 100.0 * (1 - n / expected))

    def stop(self):
        try:
            self.sniffer.stop()
        except Exception:
            pass


CLOUD_HINTS = ("googleusercontent", "amazonaws", "azure", "cloudapp", "gcp", "konami",
               "akamai", "linode", "vultr", "ovh")


def geo_lookup(ip=""):
    """Nazione/citta'/provider di un IP (vuoto = il tuo IP pubblico) via ip-api.com.
    Restituisce {} se non risponde in tempo."""
    url = (f"http://ip-api.com/json/{ip}?lang={'it' if LANG == 'it' else 'en'}"
           "&fields=status,country,countryCode,city,isp,org,hosting")
    try:
        with urllib.request.urlopen(url, timeout=1.2) as r:
            d = json.load(r)
        return d if d.get("status") == "success" else {}
    except Exception:
        return {}


def guess_platform(ttl):
    """Ipotesi sul sistema dell'avversario dal TTL dei pacchetti (solo P2P).
    Ogni sistema parte da un valore fisso e ogni router attraversato toglie 1:
    Windows/Xbox partono da 128, PlayStation/Android/iPhone da 64."""
    if ttl is None or ttl > 128:
        return "?"
    if ttl > 64:
        return T("probabile PC/Xbox", "likely PC/Xbox")
    return T("probabile PlayStation/mobile", "likely PlayStation/mobile")


def _start(fn, default):
    """Avvia fn in un thread; il risultato finisce in box[0]."""
    box = [default]

    def run():
        box[0] = fn()
    t = threading.Thread(target=run, daemon=True)
    t.start()
    return t, box


class MatchAnnouncer:
    """Appena parte una partita dice in circa 1 s che server e', dove si trova,
    quanto ping ha e suona:
      1 bip acuto = buona, 2 bip = cosi' cosi', 3 bip gravi = scarsa,
      + 1 bip lungo se l'avversario (P2P) o il server e' in un'altra nazione."""

    def __init__(self, ping_fn, beep=True):
        self.ping_fn, self.beep = ping_fn, beep
        self.matches = []   # righe per il foglio "Partite" dell'Excel
        self.latest = {}    # ip -> riga dell'ultima partita con quel server
        self.info = {}      # ip -> (tipo, dove, provider, altra nazione) gia' noti
        self.home = {}
        threading.Thread(target=self._load_home, daemon=True).start()

    def _load_home(self):
        self.home = geo_lookup()

    def __call__(self, ip, port, ttl=None):
        t0 = time.monotonic()
        # geolocalizzazione, nome host e ping partono tutti insieme
        jobs = {
            "geo": _start(lambda: geo_lookup(ip), {}),
            "host": _start(lambda: _rdns(ip), ""),
            "ping": _start(
                lambda: [self.ping_fn(ip, 300) for _ in range(2)], []),
        }
        # si aspetta al massimo 1,2 s; il nome host serve solo se manca la geolocalizzazione
        for k in ("geo", "ping"):
            jobs[k][0].join(max(0.0, 1.2 - (time.monotonic() - t0)))
        if not jobs["geo"][1][0]:
            jobs["host"][0].join(max(0.0, 1.2 - (time.monotonic() - t0)))
        geo, host, rtts = (jobs[k][1][0] for k in ("geo", "host", "ping"))

        cloud = geo.get("hosting") or any(h in (host or "").lower() for h in CLOUD_HINTS)
        kind = KIND_SERVER if cloud else KIND_P2P
        ok = [r for r in rtts if r is not None]
        ping = statistics.fmean(ok) if ok else None
        if ping is None:
            verdict, beeps, col = CONN_NA, 2, YELLOW
        elif ping <= 40 and len(ok) == len(rtts):
            verdict, beeps, col = CONN_GOOD, 1, GREEN
        elif ping <= 80:
            verdict, beeps, col = CONN_SOSO, 2, YELLOW
        else:
            verdict, beeps, col = CONN_POOR, 3, RED

        where = ", ".join(x for x in (geo.get("city"), geo.get("country")) if x) or "?"
        abroad = bool(geo.get("countryCode") and self.home.get("countryCode")
                      and geo["countryCode"] != self.home["countryCode"])
        who = "server" if cloud else T("avversario", "opponent")
        platform_guess = "" if cloud else guess_platform(ttl)
        if not geo:
            nation = T("nazione non disponibile", "country not available")
        elif abroad:
            nation = T(f"{who.upper()} IN UN'ALTRA NAZIONE ({geo.get('country')})",
                       f"{who.upper()} IN ANOTHER COUNTRY ({geo.get('country')})")
        else:
            nation = T(f"{who} nella tua stessa nazione", f"{who} in your own country")
        now = dt.datetime.now()
        secs = time.monotonic() - t0
        print(f"\n{col}{now:%H:%M:%S} >>> {T('PARTITA IN ARRIVO', 'MATCH INCOMING')} "
              f"({secs:.1f} s)  {kind}  {ip}:{port}  {where}  ping {fmt(ping, ' ms')}  "
              f"-> {T('connessione', 'connection')} {verdict}"
              f"{RESET}\n    {RED if abroad else CYAN}{nation}"
              f"{'  - ' + geo.get('isp', '') if geo.get('isp') else ''}{RESET}")
        if cloud:
            print(f"    {DIM}" + T("(partita via server: l'IP dell'avversario non si vede)",
                                   "(match via server: the opponent's IP is hidden)") + RESET)
        elif platform_guess != "?":
            print(f"    {CYAN}{T('avversario', 'opponent')}: {platform_guess} "
                  f"{DIM}(TTL {ttl}){RESET}")
        print()
        if self.beep and IS_WINDOWS:
            try:
                import winsound
                freq = 1200 if beeps == 1 else 700 if beeps == 2 else 400
                for _ in range(beeps):
                    winsound.Beep(freq, 150)
                    time.sleep(0.06)
                if abroad:
                    time.sleep(0.15)
                    winsound.Beep(500, 700)
            except Exception:
                pass
        self.latest[ip] = rec = [now.strftime(T("%d/%m/%Y", "%Y-%m-%d")),
                                 now.strftime("%H:%M:%S"), f"{ip}:{port}", kind, where, geo.get("isp", ""),
                                 YES if abroad else (NO if geo else "?"),
                                 None if ping is None else round(ping), verdict,
                                 None, None, None, None, None, None, platform_guess]
        self.info[ip] = tuple(rec[3:7]) + (platform_guess,)
        self.matches.append(rec)

    def record_for(self, ip, port):
        """Riga della partita in corso con questo server (creata se l'avviso e' mancato)."""
        rec = self.latest.get(ip)
        if rec is None:
            now = dt.datetime.now()
            known = self.info.get(ip, ("?", "?", "", "?", ""))  # stesso server gia' visto
            rec = self.latest[ip] = [now.strftime(T("%d/%m/%Y", "%Y-%m-%d")),
                                     now.strftime("%H:%M:%S"),
                                     f"{ip}:{port}", *known[:4], None, CONN_NA,
                                     None, None, None, None, None, None, known[4]]
            self.matches.append(rec)
        return rec


class MatchReport:
    """Accumula le finestre di una partita e alla fine compila la sua pagella."""

    def __init__(self, record):
        self.record, self.windows = record, 0
        self.pings, self.losses, self.gaps, self.max_gap = [], [], 0, 0.0

    def add(self, m, srv):
        self.windows += 1
        self.gaps += m["gaps"]
        self.max_gap = max(self.max_gap, m["max_gap"] or 0)
        if m["loss"] is not None:
            self.losses.append(m["loss"])
        if srv and srv["avg"] is not None:
            self.pings.append(srv["avg"])

    def finish(self, window_s):
        ping = statistics.fmean(self.pings) if self.pings else None
        loss = statistics.fmean(self.losses) if self.losses else None
        bad = (loss or 0) >= 3 or self.gaps >= 10 or (ping or 0) >= 100
        warn = (loss or 0) >= 0.5 or self.gaps >= 3 or (ping or 0) >= 60
        verdict = BAD if bad else WARN if warn else OK
        self.record[9:15] = [round(self.windows * window_s / 60, 1),
                           None if ping is None else round(ping),
                           self.gaps, round(self.max_gap) or None,
                           None if loss is None else round(loss, 1), verdict]
        return verdict, ping, loss


def _rdns(ip):
    try:
        return socket.gethostbyaddr(ip)[0]
    except OSError:
        return ""


# --------------------------------------------------------------------------
# Main loop
# --------------------------------------------------------------------------
# Colonne del file della sessione (l'ordine conta: l'Excel le legge per posizione)
_CSV_FIELDS = [
    ("Data", "Date"), ("Ora", "Time"), ("Giudizio", "Verdict"), ("Cosa non va", "What's wrong"),
    ("Segnale Wi-Fi (%)", "Wi-Fi signal (%)"), ("Segnale Wi-Fi (dBm)", "Wi-Fi signal (dBm)"),
    ("Banda Wi-Fi", "Wi-Fi band"), ("Canale Wi-Fi", "Wi-Fi channel"),
    ("Ping router (ms)", "Router ping (ms)"), ("Sbalzi ping router (ms)", "Router jitter (ms)"),
    ("Pacchetti persi router (%)", "Router packet loss (%)"),
    ("Ping internet (ms)", "Internet ping (ms)"),
    ("Ping internet peggiore (ms)", "Worst internet ping (ms)"),
    ("Sbalzi ping internet (ms)", "Internet jitter (ms)"),
    ("Pacchetti persi internet (%)", "Internet packet loss (%)"),
    ("Server partita", "Match server"),
    ("Pacchetti partita ricevuti al secondo", "Match packets received per second"),
    ("Pacchetti partita inviati al secondo", "Match packets sent per second"),
    ("Pacchetti persi in partita (%)", "Match packet loss (%)"),
    ("Scatti in partita", "Match stutters"),
    ("Pausa piu lunga in partita (ms)", "Longest match pause (ms)"),
    ("Ping server partita (ms)", "Match server ping (ms)"),
]


def diagnose(wi, gwr, net, m, srv, roamed):
    """Traduce i numeri in un giudizio e in una frase comprensibile."""
    issues = []  # (gravita', frase)

    def check(value, warn, bad, text):
        if value is None:
            return
        if value >= bad:
            issues.append((2, text.format(v=value)))
        elif value >= warn:
            issues.append((1, text.format(v=value)))

    sig = wi.get("signal")
    if sig is not None:
        check(100 - sig, 30, 50, T("segnale Wi-Fi debole ({0}%)", "weak Wi-Fi signal ({0}%)").format(sig))
    if roamed:
        issues.append((2, T("il PC ha cambiato antenna/access point",
                              "the PC switched access point")))
    if gwr:
        if gwr["sent"] and gwr["avg"] is None:
            issues.append((2, T("il router non risponde", "the router is not responding")))
        check(gwr["loss"], 0.1, 2, T("il Wi-Fi perde pacchetti ({v:.0f}%)",
                                      "Wi-Fi is losing packets ({v:.0f}%)"))
        check(gwr["jitter"], 6, 15, T("Wi-Fi instabile, il ping al router balla di {v:.0f} ms",
                                       "unstable Wi-Fi, router ping swings by {v:.0f} ms"))
        check(gwr["avg"], 10, 30, T("Wi-Fi lento, ping al router {v:.0f} ms",
                                     "slow Wi-Fi, router ping {v:.0f} ms"))
    if net["sent"] and net["avg"] is None:
        issues.append((2, T("internet non risponde", "internet is not responding")))
    else:
        check(net["loss"], 0.1, 2, T("la linea internet perde pacchetti ({v:.0f}%)",
                                      "the internet line is losing packets ({v:.0f}%)"))
        check(net["jitter"], 8, 20, T("linea instabile, il ping internet balla di {v:.0f} ms",
                                       "unstable line, internet ping swings by {v:.0f} ms"))
        check(net["avg"], 50, 100, T("ping internet alto ({v:.0f} ms)", "high internet ping ({v:.0f} ms)"))
    if m:
        check(m["gaps"], 1, 3, T("{v} scatti in partita (pause oltre 150 ms)",
                                  "{v} match stutters (pauses over 150 ms)"))
        check(m["loss"], 1, 5, T("persi {v:.0f}% dei pacchetti dal server",
                                  "lost {v:.0f}% of packets from the server"))
    if srv and not srv["no_icmp"]:
        check(srv["avg"], 70, 120, T("server della partita lontano ({v:.0f} ms)",
                                      "match server far away ({v:.0f} ms)"))

    if not issues:
        return OK, T("tutto a posto", "all good")
    level = BAD if any(g == 2 for g, _ in issues) else WARN
    issues.sort(key=lambda x: -x[0])
    return level, "; ".join(t for _, t in issues)


def num(v, nd=0):
    """Numero per il file: con la virgola in italiano, col punto in inglese."""
    if v is None:
        return ""
    return f"{v:.{nd}f}".replace(".", ",") if LANG == "it" else f"{v:.{nd}f}"


def fmt(v, suffix="", nd=0):
    return "-" if v is None else f"{v:.{nd}f}{suffix}"


def fmt_ping(label, r, good_ms, bad_ms):
    if r["no_icmp"]:
        return f"{label} {DIM}no ping{RESET}"
    ms = color(r["avg"], fmt(r["avg"], "ms"), good_ms, bad_ms)
    jit = color(r["jitter"], fmt(r["jitter"], "", 0), good_ms / 2, bad_ms / 2)
    loss = color(r["loss"], fmt(r["loss"], "%"), 0, 2)
    return f"{label} {ms} ±{jit} {loss}"


# --------------------------------------------------------------------------
# Excel colorato (serve: pip install openpyxl)
# --------------------------------------------------------------------------
# numero di colonna -> (soglia giallo, soglia rosso, True se "piu' alto = peggio")
XLSX_THRESHOLDS = {
    4: (70, 50, False),      # segnale Wi-Fi %
    8: (10, 30, True),       # ping router
    9: (6, 15, True),        # sbalzi router
    10: (0.1, 2, True),      # persi router
    11: (50, 100, True),     # ping internet
    12: (100, 200, True),    # ping internet peggiore
    13: (8, 20, True),       # sbalzi internet
    14: (0.1, 2, True),      # persi internet
    18: (1, 5, True),        # persi in partita
    19: (1, 3, True),        # scatti
    20: (150, 300, True),    # pausa piu' lunga
    21: (70, 120, True),     # ping server
}
SUMMARY_COLUMNS = (4, 8, 9, 10, 11, 13, 14, 18, 19, 21)


HISTORY_FILE = "efootball_match_history.xlsx"
_HISTORY_HEADER = [("Data", "Date"), ("Ora", "Time"), ("Server", "Server"), ("Tipo", "Type"),
                   ("Dove", "Where"), ("Provider", "Provider"),
                   ("Altra nazione", "Other country"), ("Ping iniziale (ms)", "Start ping (ms)"),
                   ("Connessione all'avvio", "Connection at start"),
                   ("Durata (min)", "Duration (min)"),
                   ("Ping medio in partita (ms)", "Avg match ping (ms)"),
                   ("Scatti", "Stutters"), ("Pausa max (ms)", "Longest pause (ms)"),
                   ("Pacchetti persi (%)", "Packet loss (%)"),
                   ("Pagella partita", "Match report"),
                   ("Avversario (P2P)", "Opponent (P2P)")]
_MATCH_WIDTHS = (11, 9, 22, 28, 24, 26, 13, 12, 16, 10, 14, 9, 11, 12, 16, 26)
_LEVEL_COLORS = (("C6EFCE", "006100"), ("FFEB9C", "9C5700"), ("FFC7CE", "9C0006"))


def write_match_rows(ws, matches, new_sheet):
    """Scrive (o aggiunge) le righe delle partite con intestazione e colori."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if new_sheet or ws.max_row < 1 or ws.cell(1, 1).value not in ("Data", "Date") \
            or ws.max_column < len(HISTORY_HEADER):
        for i, h in enumerate(HISTORY_HEADER, 1):   # crea o aggiorna l'intestazione
            c = ws.cell(1, i, h)
            c.fill = PatternFill("solid", start_color="1F4E78")
            c.font = Font(bold=True, color="FFFFFF")
        for i, w in enumerate(_MATCH_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    for m in matches:
        ws.append(m)
        r = ws.max_row
        for col in (9, 15):
            lv = (CONN_LEVEL if col == 9 else LEVEL_OF).get(m[col - 1]) \
                if col <= len(m) else None
            if lv is not None:
                bg, fg = _LEVEL_COLORS[lv]
                ws.cell(r, col).fill = PatternFill("solid", start_color=bg)
                ws.cell(r, col).font = Font(color=fg, bold=(col == 15))
        if m[6] in ("SI", "YES"):
            ws.cell(r, 7).fill = PatternFill("solid", start_color="FFC7CE")
            ws.cell(r, 7).font = Font(color="9C0006")
    ws.auto_filter.ref = ws.dimensions


def append_history(matches, folder):
    """Aggiunge le partite di questa sessione al file storico (uno solo per sempre)."""
    if not matches:
        return None
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return None
    path = os.path.join(folder, HISTORY_FILE)
    if os.path.exists(path):
        wb = load_workbook(path)
        write_match_rows(wb.active, matches, new_sheet=False)
    else:
        wb = Workbook()
        wb.active.title = T("Partite", "Matches")
        write_match_rows(wb.active, matches, new_sheet=True)
    try:
        wb.save(path)
    except PermissionError:
        print(RED + T(f"Non riesco ad aggiornare {path}: e' aperto in Excel? Le partite di "
                      f"questa sessione restano comunque nel file della sessione.",
                      f"Can't update {path}: is it open in Excel? This session's matches are "
                      f"still in the session file.") + RESET)
        return None
    return path


def make_xlsx(csv_path, matches=None):
    """Crea accanto al CSV un file .xlsx formattato. Restituisce il percorso o None."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(YELLOW + T("Per il file Excel colorato esegui 'pip install openpyxl' e poi: ",
                         "For the color-coded Excel file run 'pip install openpyxl' and then: ")
              + f"python efootball_monitor.py --excel \"{csv_path}\"{RESET}")
        return None

    try:
        with open(csv_path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh, delimiter=";"))
    except OSError as e:
        print(f"{RED}{T('Non riesco a leggere', 'Cannot read')} {csv_path}: {e}{RESET}")
        return None
    if len(rows) < 2:
        return None
    header, data = rows[0], rows[1:]

    fills = {k: PatternFill("solid", start_color=c) for k, c in
             (("g", "C6EFCE"), ("y", "FFEB9C"), ("r", "FFC7CE"), ("h", "1F4E78"))}
    fonts = {"g": Font(color="006100"), "y": Font(color="9C5700"), "r": Font(color="9C0006")}
    level_key = ("g", "y", "r")

    def to_num(v):
        try:
            return float(v.replace(",", "."))
        except (ValueError, AttributeError):
            return v if v != "" else None

    wb = Workbook()
    ws = wb.active
    ws.title = T("Andamento", "Timeline")
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = fills["h"], Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 45

    counts = defaultdict(int)
    for r in data:
        values = [to_num(v) if i >= 4 else v for i, v in enumerate(r)]
        ws.append(values)
        row = ws.max_row
        lv = LEVEL_OF.get(r[2])
        counts[lv] += 1
        k = level_key[lv] if lv is not None else None
        if k:
            for col in (3, 4):
                cell = ws.cell(row, col)
                cell.fill, cell.font = fills[k], Font(bold=(col == 3), color=fonts[k].color)
        for i in range(len(header)):
            th = XLSX_THRESHOLDS.get(i)
            v = values[i] if i < len(values) else None
            if not th or not isinstance(v, float):
                continue
            warn, bad, higher_worse = th
            x = v if higher_worse else -v
            w, b = (warn, bad) if higher_worse else (-warn, -bad)
            key = "r" if x >= b else "y" if x >= w else "g"
            cell = ws.cell(row, i + 1)
            cell.fill, cell.font = fills[key], fonts[key]

    widths = {1: 11, 3: 13, 4: 60, 16: 22}
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 12)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    # Foglio di riepilogo
    rs = wb.create_sheet(T("Riepilogo", "Summary"), 0)
    rs.column_dimensions["A"].width = 34
    rs.column_dimensions["B"].width = 16
    rs["A1"] = T("Riepilogo connessione eFootball", "eFootball connection summary")
    rs["A1"].font = Font(bold=True, size=14)
    rs["A2"] = T(f"{data[0][0]} dalle {data[0][1]} alle {data[-1][1]}",
                 f"{data[0][0]} from {data[0][1]} to {data[-1][1]}")
    rs.append([])
    rs.append([T("Giudizio", "Verdict"), T("% del tempo", "% of time")])
    for c in rs[rs.max_row]:
        c.font = Font(bold=True)
    tot = len(data)
    for lv, label in enumerate((OK, WARN, BAD)):
        rs.append([label, round(100 * counts[lv] / tot, 1)])
        k = level_key[lv]
        for c in rs[rs.max_row]:
            c.fill, c.font = fills[k], fonts[k]
    rs.append([])
    rs.append([T("Media della sessione", "Session average"), T("valore", "value")])
    for c in rs[rs.max_row]:
        c.font = Font(bold=True)
    for i in SUMMARY_COLUMNS:
        name = header[i] if i < len(header) else ""
        vals = [to_num(r[i]) for r in data if i < len(r)]
        vals = [v for v in vals if isinstance(v, float)]
        if vals:
            rs.append([name, round(statistics.fmean(vals), 1)])
    problems = defaultdict(int)
    for r in data:
        if LEVEL_OF.get(r[2], 0) != 0:
            for p in r[3].split("; "):
                key = re.sub(r"\(.*?\)", "", p.split(",")[0])  # senza i numeri
                problems[re.sub(r"^\d+\s*", "", key).strip()] += 1
    if problems:
        rs.append([])
        rs.append([T("Problemi piu frequenti", "Most frequent problems"),
                   T("quante volte", "how many times")])
        for c in rs[rs.max_row]:
            c.font = Font(bold=True)
        for p, n in sorted(problems.items(), key=lambda x: -x[1])[:6]:
            rs.append([p, n])

    if matches:
        write_match_rows(wb.create_sheet(T("Partite", "Matches"), 1), matches, new_sheet=True)

    xlsx_path = os.path.splitext(csv_path)[0] + ".xlsx"
    try:
        wb.save(xlsx_path)
    except PermissionError:
        print(RED + T(f"Non riesco a salvare {xlsx_path}: e' aperto in Excel? Chiudilo e "
                      f"riprova con --excel.",
                      f"Can't save {xlsx_path}: is it open in Excel? Close it and try again "
                      f"with --excel.") + RESET)
        return None
    return xlsx_path


set_lang("it")  # valori di default; main() li reimposta con la lingua scelta


def main():
    pre = argparse.ArgumentParser(add_help=False)
    pre.add_argument("--lang", choices=("it", "en"))
    set_lang(pre.parse_known_args()[0].lang or detect_lang())

    ap = argparse.ArgumentParser(description=T(
        "Monitor connessione per eFootball (Wi-Fi, Windows)",
        "Connection monitor for eFootball (Wi-Fi, Windows)"))
    ap.add_argument("--lang", choices=("it", "en"),
                    help=T("lingua: it o en (default: quella di Windows)",
                           "language: it or en (default: Windows language)"))
    ap.add_argument("--internet", default="1.1.1.1",
                    help=T("host internet da pingare (default 1.1.1.1)",
                           "internet host to ping (default 1.1.1.1)"))
    ap.add_argument("--gateway", help=T("IP del router (default: rilevato da solo)",
                                        "router IP (default: detected automatically)"))
    ap.add_argument("--interval", type=float, default=0.5,
                    help=T("secondi tra un ping e l'altro (default 0.5)",
                           "seconds between pings (default 0.5)"))
    ap.add_argument("--window", type=float, default=5,
                    help=T("secondi per ogni riga (default 5)",
                           "seconds per output line (default 5)"))
    ap.add_argument("--csv", default=None,
                    help=T("nome del file della sessione (default efootball_log_AAAAMMGG_HHMMSS)",
                           "session file name (default efootball_log_YYYYMMDD_HHMMSS)"))
    ap.add_argument("--sniff", action="store_true",
                    help=T("analizza il traffico UDP della partita (serve Npcap + scapy, "
                           "avvia il prompt come amministratore)",
                           "analyse the match UDP traffic (needs Npcap + scapy, "
                           "run the prompt as administrator)"))
    ap.add_argument("--muto", "--mute", dest="muto", action="store_true",
                    help=T("niente bip quando viene trovata una partita",
                           "no beeps when a match is found"))
    ap.add_argument("--iface", help=T("nome interfaccia per --sniff (default: automatica)",
                                      "network interface for --sniff (default: automatic)"))
    ap.add_argument("--excel", metavar="FILE_CSV",
                    help=T("crea solo il file Excel colorato da un CSV gia' registrato",
                           "only build the color-coded Excel file from a recorded CSV"))
    args = ap.parse_args()
    if FROZEN and len(sys.argv) == 1:
        args.sniff = True  # .exe avviato con doppio clic: analisi partita attiva

    if args.excel:
        out = make_xlsx(args.excel)
        if out:
            print(f"{T('Excel salvato in', 'Excel saved to')} {os.path.abspath(out)}")
        return

    ping_fn = make_ping_fn()
    my_ip = local_ip()
    gw = args.gateway or default_gateway()
    csv_path = args.csv or dt.datetime.now().strftime("efootball_log_%Y%m%d_%H%M%S.csv")
    timeout_ms = 1000

    print(f"{CYAN}eFootball Wi-Fi Monitor{RESET}  {T('IP locale', 'local IP')} {my_ip}  "
          f"router {gw}  internet {args.internet}")
    wi = wifi_info()
    if wi:
        print(f"Wi-Fi: {wi.get('ssid', '?')}  {wi.get('radio', '')} {wi.get('band', '')} "
              f"{T('canale', 'channel')} {wi.get('channel', '?')}")
        if wi.get("band", "").startswith("2.4") or (wi.get("channel") or 99) <= 14:
            print(YELLOW + T("Sei sulla banda 2.4 GHz: se il router lo permette usa la 5 GHz, "
                             "di solito ha meno interferenze e jitter.",
                             "You are on the 2.4 GHz band: if your router allows it use "
                             "5 GHz, it usually has less interference and jitter.") + RESET)
    elif IS_WINDOWS:
        print(YELLOW + T("Nessuna rete Wi-Fi rilevata (sei via cavo?).",
                         "No Wi-Fi network detected (are you on a cable?).") + RESET)

    targets = {"net": PingTarget("internet", args.internet, args.interval, timeout_ms, ping_fn)}
    if gw:
        targets["gw"] = PingTarget("router", gw, args.interval, timeout_ms, ping_fn)
    else:
        print(YELLOW + T("Router non rilevato: usa --gateway 192.168.1.1 (o simile).",
                         "Router not detected: use --gateway 192.168.1.1 (or similar).") + RESET)

    sniffer = None
    announcer = None
    if args.sniff:
        try:
            announcer = MatchAnnouncer(ping_fn, beep=not args.muto)
            sniffer = MatchSniffer(my_ip, iface=args.iface, on_new_remote=announcer)
            targets["srv"] = PingTarget("server", None, args.interval, timeout_ms, ping_fn,
                                          may_block_icmp=True)
            print(T("Analisi traffico partita attiva: avvia una partita online.",
                    "Match traffic analysis on: start an online match."))
        except ImportError:
            print(RED + T("scapy non installato: esegui 'pip install scapy' "
                          "(e installa Npcap da https://npcap.com).",
                          "scapy is not installed: run 'pip install scapy' "
                          "(and install Npcap from https://npcap.com).") + RESET)
        except Exception as e:
            print(RED + T(f"Sniffer non avviato ({e}). Serve Npcap e il prompt come "
                          f"amministratore. Continuo senza.",
                          f"Sniffer not started ({e}). Npcap and an administrator prompt "
                          f"are needed. Continuing without it.") + RESET)

    try:
        import openpyxl  # noqa: F401
        print(T("Il file Excel viene creato quando premi Ctrl+C: ",
                "The Excel file is created when you press Ctrl+C: ")
              + os.path.abspath(os.path.splitext(csv_path)[0] + ".xlsx"))
    except ImportError:
        print(YELLOW + T(f"openpyxl non installato: salvo solo il CSV "
                         f"{os.path.abspath(csv_path)}. Per l'Excel esegui 'pip install openpyxl'.",
                         f"openpyxl is not installed: saving only the CSV "
                         f"{os.path.abspath(csv_path)}. For Excel run 'pip install openpyxl'.")
              + RESET)
    print(DIM + T("Colonne: ping medio ±jitter perdita. Ctrl+C per fermare.",
                  "Columns: avg ping ±jitter loss. Ctrl+C to stop.") + RESET + "\n")

    f = open(csv_path, "w", newline="", encoding="utf-8-sig")  # -sig: accenti ok in Excel
    writer = csv.DictWriter(f, fieldnames=CSV_FIELDS, delimiter=";")
    writer.writeheader()

    last_bssid = wi.get("bssid")
    worst = {"gw_jitter": 0.0, "net_max": 0.0, "gaps": 0, "signal_min": None}
    verdicts = defaultdict(int)
    report = None  # pagella della partita in corso

    def close_report():
        nonlocal report
        if report:
            verdict, ping, loss = report.finish(args.window)
            col = RED if verdict == BAD else YELLOW if verdict == WARN else GREEN
            print(f"{col}{dt.datetime.now():%H:%M:%S} "
                  f"{T('Fine partita: pagella', 'Match over: report')} {verdict}  "
                  f"ping {fmt(ping, ' ms')}  {T('scatti', 'stutters')} {report.gaps}  "
                  f"{T('persi', 'lost')} {fmt(loss, '%', 1)}{RESET}")
            if announcer:
                ip = report.record[2].split(":")[0]
                if announcer.latest.get(ip) is report.record:
                    del announcer.latest[ip]  # la prossima partita avra' una riga nuova
            report = None
    try:
        while True:
            time.sleep(args.window)
            now = dt.datetime.now()
            wi = wifi_info()
            res = {k: t.take_window() for k, t in targets.items()}
            m = sniffer.take_window(args.window) if sniffer else None

            if m and "srv" in targets:
                targets["srv"].set_ip(m["ip"])
                if m["new"]:
                    print(f"{CYAN}{now:%H:%M:%S} {T('Nuova partita / server:', 'New match / server:')} "
                          f"{m['ip']}:{m['port']}{RESET}")
            elif sniffer and "srv" in targets:
                targets["srv"].set_ip(None)

            roamed = bool(wi.get("bssid") and last_bssid and wi["bssid"] != last_bssid)
            if roamed:
                print(f"{RED}{now:%H:%M:%S} {T('Cambio access point', 'Access point change')} "
                      f"(roaming): {last_bssid} -> {wi['bssid']}{RESET}")
            last_bssid = wi.get("bssid") or last_bssid

            parts = [f"{now:%H:%M:%S}"]
            if wi:
                sig = wi.get("signal")
                sig_txt = f"{sig}%" + (f"/{wi['rssi']}dBm" if "rssi" in wi else "")
                parts.append("wifi " + color(None if sig is None else 100 - sig,
                                             sig_txt, 25, 50))
                if sig is not None:
                    worst["signal_min"] = sig if worst["signal_min"] is None \
                        else min(worst["signal_min"], sig)
            if "gw" in res:
                parts.append(fmt_ping("router", res["gw"], 5, 30))
                worst["gw_jitter"] = max(worst["gw_jitter"], res["gw"]["jitter"] or 0)
            parts.append(fmt_ping("internet", res["net"], 40, 100))
            worst["net_max"] = max(worst["net_max"], res["net"]["max"] or 0)
            if m:
                gap_txt = color(m["gaps"], f"{m['gaps']} {T('scatti', 'stutters')}", 0, 3)
                if m["loss"] is not None:
                    gap_txt += " " + color(m["loss"], f"{T('persi', 'lost')} {m['loss']:.0f}%",
                                           0.99, 5)
                parts.append(f"{T('partita', 'match')} {m['rx_pps']:.0f}/{m['tx_pps']:.0f} pps "
                             f"{gap_txt} (max {fmt(m['max_gap'], 'ms')})")
                worst["gaps"] += m["gaps"]
                if "srv" in res and res["srv"]["sent"]:
                    parts.append(fmt_ping("server", res["srv"], 60, 120))
            elif sniffer:
                parts.append(f"{DIM}{T('nessuna partita', 'no match')}{RESET}")

            gwr, net = res.get("gw"), res["net"]
            srv = res.get("srv") if m else None
            level, reason = diagnose(wi, gwr, net, m, srv, roamed)
            verdicts[level] += 1
            if announcer:
                rec = announcer.record_for(m["ip"], m["port"]) if m else None
                if report and (rec is None or rec is not report.record):
                    close_report()
                if rec is not None:
                    if report is None:
                        report = MatchReport(rec)
                    report.add(m, srv)
            if level != OK:
                parts.append((RED if level == BAD else YELLOW) + reason + RESET)
            print(" | ".join(parts))
            writer.writerow(dict(zip(CSV_FIELDS, [
                now.strftime(T("%d/%m/%Y", "%Y-%m-%d")), now.strftime("%H:%M:%S"), level, reason,
                wi.get("signal", ""), wi.get("rssi", ""), wi.get("band", ""),
                wi.get("channel", ""),
                num(gwr and gwr["avg"], 1), num(gwr and gwr["jitter"], 1),
                num(gwr and gwr["loss"]),
                num(net["avg"]), num(net["max"]), num(net["jitter"], 1), num(net["loss"]),
                f"{m['ip']}:{m['port']}" if m else "",
                num(m and m["rx_pps"]), num(m and m["tx_pps"]), num(m and m["loss"], 1),
                m["gaps"] if m else "",
                num(m and m["max_gap"]),
                T("non risponde al ping", "no ping reply") if srv and srv["no_icmp"]
                else num(srv and srv["avg"]),
            ])))
            f.flush()
    except KeyboardInterrupt:
        pass
    finally:
        close_report()
        for t in targets.values():
            t.stop()
        if sniffer:
            sniffer.stop()
        f.close()

    print(f"\n{CYAN}{T('Riepilogo sessione', 'Session summary')}{RESET}")
    for k, t in targets.items():
        if not t.all_sent:
            continue
        loss = 100.0 * t.all_lost / t.all_sent
        if t.all_rtts:
            srt = sorted(t.all_rtts)
            p95 = srt[min(len(srt) - 1, int(len(srt) * 0.95))]
            print(f"  {t.name:9s} {T('media', 'avg')} {statistics.fmean(srt):.0f} ms  "
                  f"p95 {p95:.0f} ms  max {srt[-1]:.0f} ms  {T('perdita', 'loss')} {loss:.1f}% "
                  f"({t.all_sent} ping)")
        else:
            print(f"  {t.name:9s} {T('nessuna risposta', 'no reply')} ({t.all_sent} ping)")
    if worst["signal_min"] is not None:
        print(f"  {T('segnale Wi-Fi minimo', 'lowest Wi-Fi signal')}: {worst['signal_min']}%")
    tot = sum(verdicts.values())
    if tot:
        print(f"  {T('andamento', 'timeline')}: " + ", ".join(
            f"{lv} {100 * verdicts[lv] / tot:.0f}%" for lv in (OK, WARN, BAD)))
    if sniffer:
        print(f"  {T('scatti in partita, pause', 'match stutters, pauses')} (> {GAP_MS} ms): "
              f"{worst['gaps']}")
    print(T("""
Come leggerlo:
  - router con jitter o perdita alti  -> problema del Wi-Fi di casa (distanza, interferenze, 2.4 GHz)
  - router ok ma internet male        -> problema della linea / provider
  - tutto ok ma molti scatti in partita -> server o avversario lontani/instabili""", """
How to read it:
  - router with high jitter or loss   -> home Wi-Fi problem (distance, interference, 2.4 GHz)
  - router fine but internet bad      -> line / ISP problem
  - all fine but many match stutters  -> server or opponent far away / unstable"""))
    xlsx = make_xlsx(csv_path, announcer.matches if announcer else None)
    hist = append_history(announcer.matches if announcer else None,
                          os.path.dirname(os.path.abspath(csv_path)))
    if hist:
        print(f"{GREEN}{T('Storico partite aggiornato:', 'Match history updated:')} {hist}{RESET}")
    if xlsx:
        try:
            os.remove(csv_path)  # il CSV serviva solo come salvataggio durante la sessione
        except OSError:
            pass
        print(f"{GREEN}{T('Excel salvato in', 'Excel saved to')} {os.path.abspath(xlsx)}{RESET}")
    elif sum(verdicts.values()) == 0:
        try:
            os.remove(csv_path)  # sessione troppo breve: nessun dato da salvare
        except OSError:
            pass
    else:
        print(f"{T('Dati salvati in', 'Data saved to')} {os.path.abspath(csv_path)}")


if __name__ == "__main__":
    if FROZEN:
        # .exe: i file finiscono accanto al programma e la finestra non si chiude da sola
        os.chdir(os.path.dirname(sys.executable))
        try:
            main()
        except Exception:
            import traceback
            traceback.print_exc()
        finally:
            try:
                input(T("\nPremi Invio per chiudere...", "\nPress Enter to close..."))
            except (EOFError, KeyboardInterrupt):
                pass
    else:
        main()
